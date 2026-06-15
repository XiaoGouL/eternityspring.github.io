from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT_DIR = Path(__file__).resolve().parent
PPTX_PATH = OUT_DIR / "medical_internal_auditor_training_2026.pptx"
MD_PATH = OUT_DIR / "medical_internal_auditor_training_2026_outline.md"

FONT = "Microsoft YaHei"
TITLE = RGBColor(20, 52, 80)
BLUE = RGBColor(32, 92, 145)
LIGHT_BLUE = RGBColor(229, 241, 251)
ACCENT = RGBColor(0, 145, 160)
ORANGE = RGBColor(232, 128, 38)
GREEN = RGBColor(74, 141, 92)
GRAY = RGBColor(86, 96, 110)
LIGHT_GRAY = RGBColor(244, 247, 250)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(190, 63, 54)


def rgb(value):
    return RGBColor(*value)


def set_bg(slide, color=RGBColor(250, 252, 255)):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, shape_type, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    return shape


def add_text(slide, text, x, y, w, h, size=18, color=TITLE, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=None, section=None):
    if section:
        add_text(slide, section, 0.65, 0.25, 3.2, 0.28, size=10, color=ACCENT, bold=True)
    add_text(slide, title, 0.65, 0.55, 12.0, 0.6, size=28, color=TITLE, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.68, 1.16, 11.8, 0.42, size=13, color=GRAY)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.65, 1.58, 1.1, 0.05, ACCENT)


def add_footer(slide, page):
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 7.22, 13.33, 0.28, RGBColor(236, 242, 248))
    add_text(slide, "医疗体系内审员培训转训 | 供内部培训使用", 0.62, 7.25, 6.8, 0.18, size=8, color=GRAY)
    add_text(slide, f"{page:02d}", 12.25, 7.25, 0.45, 0.18, size=8, color=GRAY, align=PP_ALIGN.RIGHT)


def add_bullets(slide, bullets, x, y, w, h, size=17, color=TITLE, gap=0.28):
    y_pos = y
    for item in bullets:
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        bullet = "-" if level == 0 else "  -"
        add_text(slide, f"{bullet} {text}", x + 0.25 * level, y_pos, w - 0.25 * level, 0.28, size=size - level, color=color)
        y_pos += gap + 0.06 * max(len(text) // 36, 0)


def add_card(slide, title, bullets, x, y, w, h, accent=BLUE):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, WHITE, line=RGBColor(216, 226, 236))
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, 0.1, accent)
    add_text(slide, title, x + 0.25, y + 0.18, w - 0.5, 0.35, size=16, color=accent, bold=True)
    add_bullets(slide, bullets, x + 0.25, y + 0.68, w - 0.45, h - 0.8, size=13, color=TITLE, gap=0.25)


def add_section_slide(prs, title, subtitle, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, TITLE)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.33, 7.5, TITLE)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 5.65, 13.33, 1.85, BLUE)
    add_text(slide, title, 0.85, 2.05, 11.7, 0.75, size=36, color=WHITE, bold=True)
    add_text(slide, subtitle, 0.88, 2.95, 10.8, 0.55, size=16, color=RGBColor(221, 236, 248))
    add_text(slide, f"{page:02d}", 11.9, 6.65, 0.5, 0.22, size=9, color=RGBColor(221, 236, 248))
    return slide


def add_process(slide, steps, x=0.9, y=2.1):
    width = 1.7
    for idx, (name, desc) in enumerate(steps):
        xpos = x + idx * 2.02
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, xpos, y, width, 0.9, LIGHT_BLUE, line=RGBColor(194, 216, 236))
        add_text(slide, f"{idx + 1}", xpos + 0.08, y + 0.08, 0.28, 0.25, size=10, color=BLUE, bold=True)
        add_text(slide, name, xpos + 0.28, y + 0.18, width - 0.42, 0.3, size=15, color=TITLE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, xpos + 0.15, y + 0.54, width - 0.3, 0.25, size=9, color=GRAY, align=PP_ALIGN.CENTER)
        if idx < len(steps) - 1:
            add_text(slide, ">", xpos + width + 0.12, y + 0.26, 0.25, 0.25, size=18, color=ACCENT, bold=True)


def add_table(slide, data, x, y, w, h, header_fill=BLUE, font_size=10):
    rows, cols = len(data), len(data[0])
    table_shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = data[r][c]
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            fill = cell.fill
            fill.solid()
            fill.fore_color.rgb = header_fill if r == 0 else (WHITE if r % 2 else LIGHT_GRAY)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
                for run in paragraph.runs:
                    run.font.name = FONT
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = WHITE if r == 0 else TITLE
                    run.font.bold = r == 0
    return table_shape


slides_meta = []


def remember(title, bullets):
    slides_meta.append((title, bullets))


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1
    slide = prs.slides.add_slide(blank)
    set_bg(slide, RGBColor(248, 252, 255))
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.33, 1.55, TITLE)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 6.75, 13.33, 0.75, BLUE)
    add_text(slide, "医疗体系内审员培训转训", 0.82, 1.95, 11.6, 0.78, size=38, color=TITLE, bold=True)
    add_text(slide, "法规最新要求 | 质量管理体系 | ISO 13485 | 采购与供应商控制", 0.86, 2.86, 11.2, 0.42, size=18, color=ACCENT, bold=True)
    add_text(slide, "适用对象：采购、质量、仓储、生产、研发及供应商管理相关人员", 0.88, 3.58, 10.4, 0.34, size=15, color=GRAY)
    add_text(slide, "基于深圳医疗协会医疗体系内审员培训内容整理", 0.88, 4.05, 10.4, 0.34, size=13, color=GRAY)
    add_text(slide, "2026 内部培训版", 0.88, 5.15, 4.0, 0.3, size=14, color=BLUE, bold=True)
    add_footer(slide, 1)
    remember("医疗体系内审员培训转训", [
        "建议作为 60-90 分钟内部培训材料使用。",
        "培训重点面向采购工程师和供应商管理岗位。",
    ])

    # 2
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "本次培训目标", "把外部内审员培训转化为内部可执行的采购与质量管理动作", "培训导入")
    add_card(slide, "听完后应能做到", [
        "说清医疗器械监管框架和近期变化",
        "理解质量管理体系的过程方法和风险思维",
        "掌握 ISO 13485 对采购控制的核心要求",
        "能用内审视角识别采购和供应商管理缺口",
    ], 0.75, 2.0, 5.75, 3.45, BLUE)
    add_card(slide, "建议互动方式", [
        "每个模块结合 1 个公司实际案例",
        "现场抽查 1 份供应商档案或采购记录",
        "结束前形成部门整改/优化清单",
    ], 6.85, 2.0, 5.75, 3.45, ACCENT)
    add_footer(slide, 2)
    remember("本次培训目标", ["把法规、体系和标准要求转换成采购端可执行事项。"])

    # 3
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "为什么采购工程师要懂内审", "采购不是简单下单，而是医疗器械全生命周期质量控制的前置关口", "角色认知")
    add_card(slide, "采购对产品质量的影响", [
        "关键物料、外协过程和服务会直接影响安全有效",
        "供应商变更、工艺漂移、证照过期会放大合规风险",
        "采购记录是追溯、召回和监管检查的重要证据",
    ], 0.7, 1.95, 3.85, 3.7, BLUE)
    add_card(slide, "内审能帮助采购发现什么", [
        "流程是否按文件执行",
        "记录是否真实、完整、可追溯",
        "供应商准入、监控和重评是否闭环",
        "质量协议和变更通知是否有效",
    ], 4.75, 1.95, 3.85, 3.7, ACCENT)
    add_card(slide, "采购人员的底线思维", [
        "不从未批准供应商采购关键物料",
        "不接受资质不清或证据缺失的产品",
        "不绕过质量部门处理异常",
    ], 8.8, 1.95, 3.85, 3.7, ORANGE)
    add_footer(slide, 3)
    remember("为什么采购工程师要懂内审", ["采购端是质量风险输入端，也是供应链合规证据的第一责任链条之一。"])

    # 4
    add_section_slide(prs, "第一部分：医疗器械法规与监管变化", "从合规框架看采购、供应商和追溯要求", 4)
    remember("第一部分：医疗器械法规与监管变化", ["本部分用于建立共同法规语言。"])

    # 5
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "医疗器械法规框架一览", "培训中可按“法规—规章—规范—标准—公司文件”逐层理解", "法规框架")
    table = [
        ["层级", "示例", "对采购/内审的关注点"],
        ["法规", "《医疗器械监督管理条例》", "注册人/备案人主体责任、全生命周期监管、违法后果"],
        ["部门规章", "注册与备案、生产监督、经营监督等办法", "许可备案、变更、生产经营活动边界"],
        ["质量规范", "生产质量管理规范、经营质量管理规范", "体系运行、采购验收、仓储追溯、委托/外协控制"],
        ["强制/推荐标准", "GB/YY 标准、ISO 13485 等", "产品和过程要求、验证确认、记录和证据"],
        ["公司文件", "质量手册、程序、SOP、表单", "内审直接依据，必须与法规和标准保持一致"],
    ]
    add_table(slide, table, 0.75, 1.95, 11.85, 3.75, font_size=10)
    add_text(slide, "内审提示：外部法规不是“挂在墙上的文件”，应被转化进程序、职责、记录和日常操作。", 0.9, 6.0, 11.4, 0.35, size=15, color=ORANGE, bold=True)
    add_footer(slide, 5)
    remember("医疗器械法规框架一览", ["内审时要检查法规要求是否已经转化为公司文件和执行记录。"])

    # 6
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "近期监管变化：需要内部培训关注的重点", "以下为公开监管信息的培训化整理，具体执行以现行有效法规原文为准", "法规变化")
    add_card(slide, "经营质量管理规范", [
        "新版《医疗器械经营质量管理规范》自 2024-07-01 施行",
        "强化采购、验收、仓储、销售、售后和追溯全过程管理",
        "关注 UDI、电子证照、直调、多仓协同等新场景",
    ], 0.7, 1.9, 3.85, 3.95, BLUE)
    add_card(slide, "生产质量管理规范", [
        "国家药监局 2025 年发布新版生产质量管理规范",
        "自 2026-11-01 施行，过渡期内应完成文件和体系升级",
        "强调质量保证、验证确认、委托生产和外协加工受控",
    ], 4.75, 1.9, 3.85, 3.95, ACCENT)
    add_card(slide, "共同监管趋势", [
        "全生命周期质量责任",
        "风险管理和持续改进",
        "数据真实、完整、可追溯",
        "注册人/备案人对供应链不能“甩锅”",
    ], 8.8, 1.9, 3.85, 3.95, ORANGE)
    add_footer(slide, 6)
    remember("近期监管变化：需要内部培训关注的重点", ["监管趋势从结果合格延伸到过程受控、责任清晰和数据可追溯。"])

    # 7
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "医疗器械分类与采购合规边界", "采购前先确认“买的是什么、用于哪里、监管属性是什么”", "法规基础")
    add_card(slide, "按风险分类", [
        "第一类：常规管理可保证安全有效",
        "第二类：需严格控制管理",
        "第三类：植入、支持维持生命等较高风险，控制最严格",
    ], 0.7, 1.9, 3.85, 3.8, BLUE)
    add_card(slide, "采购前应确认", [
        "产品注册证/备案凭证及适用范围",
        "生产/经营资质、授权链条和有效期",
        "规格型号、关键技术要求、标签说明书要求",
        "是否涉及冷链、无菌、植入、UDI 或特殊储运",
    ], 4.75, 1.9, 3.85, 3.8, ACCENT)
    add_card(slide, "典型红线", [
        "证照过期或授权链不完整",
        "实际采购型号与注册/备案信息不一致",
        "供应商未按约定提前通知重大变更",
        "未按规定保留追溯记录",
    ], 8.8, 1.9, 3.85, 3.8, RED)
    add_footer(slide, 7)
    remember("医疗器械分类与采购合规边界", ["采购活动需要先判断产品监管属性，不能只依据商业报价。"])

    # 8
    add_section_slide(prs, "第二部分：质量管理体系与内审方法", "把法规要求变成过程、证据和持续改进", 8)
    remember("第二部分：质量管理体系与内审方法", ["本部分讲体系运行和内审工作方法。"])

    # 9
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "质量管理体系的过程方法", "医疗器械质量体系不是文件集合，而是一组相互作用的过程", "质量体系")
    steps = [
        ("输入", "法规/客户/风险/资源"),
        ("过程", "采购/生产/检验/仓储"),
        ("输出", "合格产品/记录/服务"),
        ("监测", "指标/审核/投诉/不良事件"),
        ("改进", "CAPA/变更/管理评审"),
    ]
    add_process(slide, steps, 1.0, 2.0)
    add_card(slide, "内审看什么", [
        "过程是否有明确负责人、输入、输出和控制方法",
        "关键风险是否被识别并设置控制点",
        "记录能否证明过程按要求执行",
        "问题是否进入纠正预防和持续改进闭环",
    ], 1.15, 4.25, 10.95, 1.55, BLUE)
    add_footer(slide, 9)
    remember("质量管理体系的过程方法", ["过程方法的核心是输入、输出、控制、证据和改进闭环。"])

    # 10
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "内审的定位与原则", "内审是自我发现问题的机制，不是为了追责而追责", "内审方法")
    add_card(slide, "定位", [
        "验证体系是否符合：法规、标准、公司文件",
        "验证体系是否有效：是否能稳定产出符合要求的结果",
        "识别风险和改进机会，为管理评审提供输入",
    ], 0.7, 1.95, 3.85, 3.6, BLUE)
    add_card(slide, "原则", [
        "基于证据：文件、记录、访谈、现场观察相互印证",
        "基于风险：优先审关键物料、关键供应商、异常高发过程",
        "独立客观：审核员不审核自己直接负责的工作",
    ], 4.75, 1.95, 3.85, 3.6, ACCENT)
    add_card(slide, "常用证据", [
        "程序文件、SOP、记录表单",
        "采购订单、验收记录、供应商档案",
        "质量协议、审计报告、CAPA/变更记录",
        "现场状态、系统数据和人员访谈",
    ], 8.8, 1.95, 3.85, 3.6, ORANGE)
    add_footer(slide, 10)
    remember("内审的定位与原则", ["内审结论要基于客观证据，关注体系符合性与有效性。"])

    # 11
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "内审流程：从计划到闭环", "好的内审不是发现问题就结束，而是验证纠正措施有效", "内审方法")
    add_process(slide, [
        ("年度计划", "范围/频次/资源"),
        ("审核准备", "法规/文件/检查表"),
        ("现场审核", "访谈/观察/抽样"),
        ("报告问题", "事实/依据/分级"),
        ("整改验证", "原因/CAPA/证据"),
        ("输入评审", "趋势/资源/改进"),
    ], 0.5, 2.0)
    add_bullets(slide, [
        "问题描述建议使用：事实 + 依据 + 风险影响 + 证据编号。",
        "纠正措施不等于“重新培训”本身，应包含原因分析、措施、责任人、期限和有效性验证。",
        "采购相关问题要关注是否影响已采购批次、库存、在制品和已交付产品。",
    ], 1.0, 4.45, 11.3, 1.15, size=15, color=TITLE, gap=0.32)
    add_footer(slide, 11)
    remember("内审流程：从计划到闭环", ["审核发现必须通过 CAPA 和有效性验证闭环。"])

    # 12
    add_section_slide(prs, "第三部分：ISO 13485 与采购控制", "用标准条款拆解供应商准入、采购信息和来料验证", 12)
    remember("第三部分：ISO 13485 与采购控制", ["本部分把 ISO 13485 落到采购端。"])

    # 13
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "ISO 13485:2016 的结构与采购相关条款", "采购既受 7.4 直接约束，也与外包、记录、追溯、变更等条款联动", "ISO 13485")
    table = [
        ["条款", "主题", "采购相关理解"],
        ["4.1.5", "外包过程控制", "外协、委托检验、第三方服务不能脱离体系控制"],
        ["4.2.4/4.2.5", "文件与记录", "供应商清单、资质、协议、订单、验收记录须受控"],
        ["7.4.1", "采购过程", "评价、选择、监控、重评供应商，且与风险相匹配"],
        ["7.4.2", "采购信息", "明确规格、验收准则、人员/设备/过程/QMS 要求及变更通知"],
        ["7.4.3", "采购产品验证", "到货验收、供应商现场验证或其他验证方式要有证据"],
        ["8.2/8.5", "监测改进/CAPA", "供应商绩效、投诉、不合格、偏差和 CAPA 闭环"],
    ]
    add_table(slide, table, 0.65, 1.85, 12.05, 4.35, font_size=9)
    add_footer(slide, 13)
    remember("ISO 13485:2016 的结构与采购相关条款", ["采购控制的主条款是 7.4，但要联动外包、记录、追溯和 CAPA。"])

    # 14
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "ISO 13485 7.4：采购控制三件事", "标准要求不只是“买合格品”，而是让采购全过程可控", "ISO 13485")
    add_card(slide, "7.4.1 采购过程", [
        "建立书面程序，确保采购产品符合规定要求",
        "建立供应商评价、选择、监控、重评准则",
        "准则考虑供应商能力、绩效、采购产品影响和医疗器械风险",
    ], 0.7, 1.9, 3.85, 3.95, BLUE)
    add_card(slide, "7.4.2 采购信息", [
        "采购要求应清晰、完整、经适当批准",
        "必要时包括产品规格、验收准则、过程/设备、人员资格、QMS 要求",
        "供应商应在影响符合性的变更实施前通知组织",
    ], 4.75, 1.9, 3.85, 3.95, ACCENT)
    add_card(slide, "7.4.3 采购产品验证", [
        "规定并实施验证活动，确认采购产品满足要求",
        "验证强度与风险、供应商评价结果匹配",
        "如在供应商现场验证，应在采购信息中明确安排和放行方法",
    ], 8.8, 1.9, 3.85, 3.95, ORANGE)
    add_footer(slide, 14)
    remember("ISO 13485 7.4：采购控制三件事", ["采购控制可概括为供应商控制、采购要求控制和采购产品验证。"])

    # 15
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "供应商与采购品风险分级", "控制强度应与风险匹配，避免所有供应商“一刀切”", "供应商管理")
    table = [
        ["等级", "对象示例", "建议控制方式"],
        ["A 关键", "直接影响安全有效的关键原材料/组件、灭菌、外协关键过程", "现场审核或等效评价；质量协议；首件/验证；绩效月度或季度监控"],
        ["B 重要", "影响质量但可通过常规检验控制的物料、关键包装、校准/检验服务", "资质审核；样品确认；定期绩效评估；必要时二方审核"],
        ["C 一般", "低风险辅料、办公服务、非关键耗材", "基础资质；采购验收；异常触发重评"],
    ]
    add_table(slide, table, 0.75, 1.95, 11.85, 2.85, font_size=9)
    add_bullets(slide, [
        "风险分级输入：产品用途、与患者/使用者接触程度、是否无菌/植入、历史质量表现、替代难度、法规要求。",
        "分级结果应驱动：准入材料、审核频次、来料检验水平、质量协议深度、绩效监控频率。",
    ], 0.9, 5.25, 11.5, 0.85, size=14, color=TITLE, gap=0.32)
    add_footer(slide, 15)
    remember("供应商与采购品风险分级", ["供应商控制要基于产品质量影响和医疗器械风险。"])

    # 16
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "供应商准入：资料与评价证据", "合格供应商清单应由证据支撑，而不是历史习惯", "供应商管理")
    add_card(slide, "资质与合法性", [
        "营业执照、生产/经营许可证或备案凭证",
        "产品注册证/备案凭证、适用范围和规格型号",
        "授权链条、进口产品代理/报关等必要证明",
    ], 0.7, 1.85, 3.85, 4.1, BLUE)
    add_card(slide, "质量能力", [
        "ISO 13485/ISO 9001 或同等 QMS 证据",
        "过程能力、检验能力、关键设备和人员资质",
        "样品确认、首件检验、试用/验证结果",
        "历史绩效、投诉、召回、不良事件配合能力",
    ], 4.75, 1.85, 3.85, 4.1, ACCENT)
    add_card(slide, "审批与维护", [
        "准入评价结论和批准人应清晰",
        "纳入合格供应商清单后方可采购",
        "证照到期、重大变更、绩效异常触发重评",
        "停用/限用供应商应及时同步采购系统",
    ], 8.8, 1.85, 3.85, 4.1, ORANGE)
    add_footer(slide, 16)
    remember("供应商准入：资料与评价证据", ["合格供应商清单必须有准入评价和持续维护证据。"])

    # 17
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "采购信息与质量协议：要把要求写清楚", "如果要求没有被传递，后续验收和争议处理就缺少依据", "采购执行")
    table = [
        ["文件/信息", "至少应覆盖的内容"],
        ["采购订单/合同", "名称、规格型号、数量、版本、图纸/标准、批号/效期、交付和储运要求"],
        ["技术/质量要求", "关键质量属性、检验标准、COA/COC、标签说明书、UDI、包装和清洁/无菌要求"],
        ["质量协议", "质量责任、变更通知、偏差和不合格处理、追溯、召回、投诉/不良事件配合"],
        ["外包/外协约定", "工艺控制、验证状态、放行权限、记录留存、监管检查配合、再委托限制"],
        ["变更通知", "影响符合性的物料、工艺、场地、设备、关键人员、检验方法、供应链变化"],
    ]
    add_table(slide, table, 0.75, 1.9, 11.85, 4.05, font_size=9)
    add_footer(slide, 17)
    remember("采购信息与质量协议：要把要求写清楚", ["采购信息越清楚，验收、变更控制和追溯越有依据。"])

    # 18
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "到货验收、追溯与记录保存", "验收不是收货签字，而是采购产品放行前的质量确认", "采购执行")
    add_card(slide, "到货验收看什么", [
        "供应商、产品、规格、数量、批号/序列号与订单一致",
        "包装、标签、说明书、合格证明、效期和运输条件",
        "冷链/温控、无菌包装、植入/高值耗材等特殊要求",
    ], 0.7, 1.9, 3.85, 3.8, BLUE)
    add_card(slide, "追溯要做到什么", [
        "能从采购记录追到供应商、批次、验收、库存、使用/销售去向",
        "UDI 适用时纳入验收、出库复核和信息系统控制",
        "异常批次能快速隔离、评估和召回/追回",
    ], 4.75, 1.9, 3.85, 3.8, ACCENT)
    add_card(slide, "记录原则", [
        "真实、准确、完整、及时、可追溯",
        "电子记录应有权限、备份、审计追踪或等效控制",
        "保存期限按法规、标准和公司文件中更严要求执行",
    ], 8.8, 1.9, 3.85, 3.8, ORANGE)
    add_footer(slide, 18)
    remember("到货验收、追溯与记录保存", ["验收和追溯记录是证明采购产品符合要求的重要证据。"])

    # 19
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "委托、外协与服务采购控制", "外部方完成的过程，仍然需要纳入组织的质量管理体系", "外部过程")
    add_card(slide, "常见外部过程", [
        "外协加工、灭菌、包装、检验检测、校准",
        "第三方仓储/物流、冷链运输、售后服务",
        "软件/信息系统服务、电子记录平台维护",
    ], 0.7, 1.95, 3.85, 3.6, BLUE)
    add_card(slide, "控制重点", [
        "委托前进行能力和质量体系评价",
        "明确质量协议、工艺/检验标准和记录要求",
        "关键过程验证确认状态持续受控",
        "再委托、变更、偏差、不合格须受控",
    ], 4.75, 1.95, 3.85, 3.6, ACCENT)
    add_card(slide, "内审追问", [
        "哪些过程外包了？风险分级是什么？",
        "有没有批准的供应商和质量协议？",
        "外包方记录如何获取、审核和保存？",
        "发生异常时谁决策放行/停用？",
    ], 8.8, 1.95, 3.85, 3.6, ORANGE)
    add_footer(slide, 19)
    remember("委托、外协与服务采购控制", ["外包不等于外责，组织仍要证明外部过程受控。"])

    # 20
    add_section_slide(prs, "第四部分：采购内审实操", "用检查表、典型问题和行动清单推动落地", 20)
    remember("第四部分：采购内审实操", ["本部分可直接用于部门自查。"])

    # 21
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "采购内审检查表：建议抽样问题", "可根据公司程序编号、产品类别和供应商等级继续细化", "内审工具")
    table = [
        ["检查主题", "审核问题", "证据示例"],
        ["供应商准入", "关键供应商是否已评价批准？评价是否基于风险？", "供应商档案、评价表、ASL"],
        ["证照资质", "供应商和产品证照是否有效、范围是否匹配？", "许可证/备案、注册证/备案、授权"],
        ["采购信息", "订单/合同是否完整传递技术和质量要求？", "订单、图纸、标准、质量协议"],
        ["变更控制", "供应商变更是否提前通知并经评估批准？", "变更通知、评估记录、验证报告"],
        ["来料验收", "验收项目、抽样、判定和不合格处理是否符合程序？", "验收记录、COA/COC、NCR"],
        ["绩效重评", "供应商绩效是否定期统计并触发措施？", "KPI、投诉/退货、CAPA、重评记录"],
    ]
    add_table(slide, table, 0.55, 1.75, 12.25, 4.85, font_size=8)
    add_footer(slide, 21)
    remember("采购内审检查表：建议抽样问题", ["检查表要连接依据、问题和客观证据。"])

    # 22
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "典型不符合与风险提示", "以下问题在内审、客户审核和监管检查中较常见", "问题识别")
    add_card(slide, "文件与清单类", [
        "合格供应商清单与系统可下单供应商不一致",
        "供应商证照过期仍未暂停采购",
        "质量协议版本旧，未覆盖变更通知和追溯要求",
    ], 0.7, 1.9, 3.85, 3.85, RED)
    add_card(slide, "执行与记录类", [
        "采购订单未写明关键规格/版本/验收要求",
        "验收记录缺批号、效期、标签/包装确认",
        "紧急采购后未补齐质量评估和批准证据",
    ], 4.75, 1.9, 3.85, 3.85, ORANGE)
    add_card(slide, "闭环与改进类", [
        "供应商绩效异常未触发 CAPA 或重评",
        "供应商变更未经内部风险评估即接收产品",
        "同类来料不合格重复发生但根因分析不足",
    ], 8.8, 1.9, 3.85, 3.85, BLUE)
    add_footer(slide, 22)
    remember("典型不符合与风险提示", ["内审发现要避免只停留在表面记录缺失，应追到流程和风险控制。"])

    # 23
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "培训落地：采购部门行动清单", "建议培训后 1 周内形成部门自查和整改计划", "行动计划")
    add_card(slide, "立即自查", [
        "抽查关键供应商档案完整性和证照有效期",
        "核对 ASL、ERP/采购系统和实际采购记录一致性",
        "抽查最近 3 批关键物料验收与追溯记录",
    ], 0.7, 1.9, 3.85, 3.8, BLUE)
    add_card(slide, "短期优化", [
        "更新供应商风险分级规则和重评频次",
        "补强质量协议模板：变更、追溯、召回、数据完整性",
        "建立证照到期和绩效异常预警",
    ], 4.75, 1.9, 3.85, 3.8, ACCENT)
    add_card(slide, "持续机制", [
        "每月 review 供应商质量 KPI",
        "重大异常纳入 CAPA 和管理评审输入",
        "采购、质量、研发、仓储联合评审关键供应商变化",
    ], 8.8, 1.9, 3.85, 3.8, GREEN)
    add_footer(slide, 23)
    remember("培训落地：采购部门行动清单", ["培训输出应转化为供应商档案、协议模板、系统权限和 KPI 的改进。"])

    # 24
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "课堂讨论 / 小测", "可用于培训结束前确认理解效果", "互动")
    add_bullets(slide, [
        "问题 1：某关键供应商 ISO 13485 证书过期 2 个月，但历史供货稳定，是否可以继续采购？需要哪些评估和批准？",
        "问题 2：供应商将关键原材料产地从 A 厂变更为 B 厂，但没有提前通知。采购、质量和研发应如何处理？",
        "问题 3：内审抽样发现 3 份来料验收记录缺少批号，这属于记录问题还是追溯风险？如何写不符合？",
        "问题 4：哪些供应商必须签质量协议？质量协议应至少覆盖哪些条款？",
        "问题 5：公司如果使用第三方冷链运输，采购和质量各自应保留哪些证据？",
    ], 0.95, 1.95, 11.5, 3.1, size=15, color=TITLE, gap=0.46)
    add_text(slide, "建议做法：每题请 1 位同事回答，再由质量/采购负责人补充公司现行流程要求。", 1.0, 5.75, 11.0, 0.35, size=14, color=ORANGE, bold=True)
    add_footer(slide, 24)
    remember("课堂讨论 / 小测", ["用实际场景验证同事是否理解供应商控制和内审证据。"])

    # 25
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "参考资料与使用提示", "培训内容建议与公司现行质量手册、程序文件和表单一起使用", "参考")
    add_bullets(slide, [
        "《医疗器械监督管理条例》及配套注册、生产、经营监管规章。",
        "国家药监局公告 2023 年第 153 号：《医疗器械经营质量管理规范》，2024-07-01 施行。",
        "国家药监局公告 2025 年第 107 号：《医疗器械生产质量管理规范》，2026-11-01 施行。",
        "ISO 13485:2016《医疗器械 质量管理体系 用于法规的要求》，重点关注 4.1.5、4.2、7.4、8.5。",
        "公司内部质量手册、采购控制程序、供应商管理程序、来料检验程序、CAPA/变更控制程序。",
    ], 0.95, 1.95, 11.5, 2.85, size=14, color=TITLE, gap=0.35)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 5.25, 11.45, 0.85, LIGHT_BLUE, line=RGBColor(194, 216, 236))
    add_text(slide, "使用提示：本 PPT 为内部培训资料，不替代法规原文、认证机构要求或公司正式受控文件；用于具体项目时，请按最新法规原文和公司程序确认。", 1.18, 5.46, 10.95, 0.32, size=13, color=TITLE, bold=True)
    add_footer(slide, 25)
    remember("参考资料与使用提示", ["培训内容应以现行有效法规和公司受控文件为最终依据。"])

    prs.save(PPTX_PATH)


def build_markdown():
    lines = [
        "# 医疗体系内审员培训转训讲稿大纲",
        "",
        "本大纲与 `medical_internal_auditor_training_2026.pptx` 配套，用于内部授课前备课和后续按公司流程微调。",
        "",
        "## 建议授课节奏",
        "",
        "- 适用对象：采购、质量、仓储、生产、研发及供应商管理相关人员。",
        "- 建议时长：60-90 分钟，可根据现场讨论深度调整。",
        "- 建议方式：法规框架讲解 + 采购案例讨论 + 抽样查看 1 份供应商档案/采购记录。",
        "",
        "## 逐页讲解提示",
        "",
    ]
    for idx, (title, bullets) in enumerate(slides_meta, start=1):
        lines.append(f"### {idx}. {title}")
        for bullet in bullets:
            lines.append(f"- {bullet}")
        lines.append("")
    lines.extend([
        "## 培训后建议输出",
        "",
        "- 关键供应商档案自查清单。",
        "- 供应商风险分级和重评频次更新建议。",
        "- 质量协议模板补充项：变更通知、追溯、召回、数据完整性、再委托限制。",
        "- 采购系统供应商状态与合格供应商清单一致性核对结果。",
        "- 对近期来料异常和供应商绩效问题的 CAPA 跟踪表。",
        "",
        "## 注意事项",
        "",
        "本资料为内部培训辅助材料，不替代法规原文、认证机构要求或公司正式受控文件。用于正式培训前，建议由质量负责人结合公司实际流程、产品类型和受控文件编号进行最终确认。",
        "",
    ])
    MD_PATH.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    build_deck()
    build_markdown()
    print(f"Generated: {PPTX_PATH}")
    print(f"Generated: {MD_PATH}")
